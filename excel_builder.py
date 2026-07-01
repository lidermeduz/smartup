"""
excel_builder.py — bitta deal'dan haqiqiy "Спецификация" Excel faylini yasaydi.

build_spec(deal, company) -> tayyor .xlsx fayl yo'lini qaytaradi.

Format haqiqiy namuna fayllarga (Спец PFL / Спец GMX) moslangan:
  Ustunlar: №, Номенклатура, ИКПУ, кол-во, Цена с НДС, Стоимость поставки,
            ставка, НДС, Стоим. поставки с учетом НДС.

НДС hisobi (MUHIM — Smartup'ning haqiqiy maydonlari bilan):
  - product_price  = Цена с НДС (НДС ICHIDA bo'lgan dona narx)
  - sold_amount    = qty * price = Стоим. поставки с учетом НДС (J)
  - vat_amount     = НДС summasi (I)
  - net = sold_amount - vat_amount = Стоимость поставки (G, НДСsiz)
  - ставка = vat_percent / 100 (masalan 0.12)
"""

import os
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

import config


def _g(d: dict, *keys, default=""):
    """Bir nechta mumkin bo'lgan kalitlardan birinchi topilganini oladi."""
    for k in keys:
        if k in d and d[k] not in (None, ""):
            return d[k]
    return default


def _num(value, default=0.0) -> float:
    """Stringni floatga aylantiradi (bo'sh bo'lsa default)."""
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _date_only(value: str) -> str:
    """ '29.06.2026 14:30:00' -> '29.06.2026' (vaqtni kesib tashlaydi). """
    if not value:
        return ""
    return str(value).split(" ")[0]


def _clean_product_name(name: str) -> str:
    """Mahsulot nomidan ishlab chiqaruvchi qo'shimchasini olib tashlaydi.
    Haqiqiy format: `<nom> / "PERFECTFOODLAB" OOO / г. Ташкент`
    -> `<nom>`.
    Ishlab chiqaruvchi qismi DOIM ' / "' (yoki ' / «') bilan boshlanadi.
    Nomning o'z ichidagi qo'shtirnoq (masalan `ПРЕПАРАТ "X" 300мг`) hech
    qachon ' / ' bilan oldinda kelmaydi, shuning uchun buzilmaydi."""
    if not name:
        return ""
    s = str(name)
    idxs = [s.find(sep) for sep in (' / "', ' / «', ' / “') if s.find(sep) != -1]
    if idxs:
        s = s[:min(idxs)]
    return s.strip().strip("/").strip()


def _ikpu_for(product_name: str) -> str:
    """Mahsulot nomidan ИКПУ (МХИК) ni aniqlaydi.
    Order export ИКПУ qaytarmagani uchun nom bo'yicha toifaga ajratamiz."""
    n = (product_name or "").lower()
    for code, keywords in config.IKPU_BY_KEYWORD.items():
        if any(kw in n for kw in keywords):
            return code
    return config.IKPU_DEFAULT


# ── Summani so'z bilan yozish (rus tilida) ──
_ONES_M = ['', 'один', 'два', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять']
_ONES_F = ['', 'одна', 'две', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять']
_TEENS = ['десять', 'одиннадцать', 'двенадцать', 'тринадцать', 'четырнадцать',
          'пятнадцать', 'шестнадцать', 'семнадцать', 'восемнадцать', 'девятнадцать']
_TENS = ['', '', 'двадцать', 'тридцать', 'сорок', 'пятьдесят', 'шестьдесят',
         'семьдесят', 'восемьдесят', 'девяносто']
_HUNDREDS = ['', 'сто', 'двести', 'триста', 'четыреста', 'пятьсот', 'шестьсот',
             'семьсот', 'восемьсот', 'девятьсот']


def _plural(n: int, one: str, few: str, many: str) -> str:
    n = abs(n) % 100
    if 11 <= n <= 14:
        return many
    d = n % 10
    if d == 1:
        return one
    if 2 <= d <= 4:
        return few
    return many


def _triple(num: int, fem: bool) -> list:
    """0..999 sonni so'zlarga (ro'yxat)."""
    w = []
    h, t, o = num // 100, (num % 100) // 10, num % 10
    if h:
        w.append(_HUNDREDS[h])
    if t == 1:
        w.append(_TEENS[o])
    else:
        if t:
            w.append(_TENS[t])
        if o:
            w.append((_ONES_F if fem else _ONES_M)[o])
    return w


def _rus_amount_words(n: float) -> str:
    """Butun sonni rus tilida so'z bilan yozadi (masalan 'один миллион ...')."""
    n = int(round(n))
    if n == 0:
        return "ноль"
    parts = []
    millions = n // 1_000_000
    thousands = (n // 1000) % 1000
    rest = n % 1000
    if millions:
        parts += _triple(millions, False)
        parts.append(_plural(millions, 'миллион', 'миллиона', 'миллионов'))
    if thousands:
        parts += _triple(thousands, True)
        parts.append(_plural(thousands, 'тысяча', 'тысячи', 'тысяч'))
    if rest:
        parts += _triple(rest, False)
    return " ".join(parts)


def _est_lines(text, width_chars: int) -> int:
    """Matn berilgan ustun kengligida necha qatorga sig'ishini taxminlaydi."""
    if not text:
        return 1
    total = 0
    for part in str(text).split("\n"):
        n = len(part)
        total += max(1, (n + width_chars - 1) // max(1, width_chars))
    return total


def _main_bank(buyer: dict, field: str) -> str:
    """Mijozning ASOSIY (is_main='Y') bank hisobidan maydonni oladi."""
    accounts = (buyer or {}).get("bank_accounts") or []
    for acc in accounts:
        if (acc.get("is_main") or "").upper() == "Y":
            return acc.get(field) or ""
    return (accounts[0].get(field) or "") if accounts else ""


def build_spec(deal: dict, company: dict, buyer: dict = None) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Спецификация"

    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    bold = Font(name="Arial", size=9, bold=True)
    reg = Font(name="Arial", size=9)
    small = Font(name="Arial", size=8)

    # A  B   C   D    E   F   G    H   I   J
    widths = [4, 22, 14, 19, 7, 12, 14, 7, 12, 15]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Sarlavha (bitta katakda 2 qator) ──
    spec_no = _g(deal, "delivery_number", default="")
    spec_date = _date_only(_g(deal, "delivery_date", "deal_time", "booked_date"))
    dog_no = _g(deal, "contract_number", default="")
    dog_date = ""  # shartnoma sanasi javobda yo'q
    ws.merge_cells("B2:J2")
    ws["B2"] = (f"Спецификация № {spec_no} от {spec_date}\n"
                f"Приложение к дог № {dog_no} от {dog_date}")
    ws["B2"].font = bold
    ws["B2"].alignment = center
    ws.row_dimensions[2].height = 32

    # ── Поставщик (chap) va Покупатель (o'ng) bloklari ──
    supplier_rows = [
        ("Поставщик:", company.get("supplier_name", company["name"])),
        ("Адрес:", company.get("supplier_address", "")),
        ("Тел:", company.get("supplier_phone", "")),
        ("ИНН:", company.get("supplier_inn", "")),
        ("Р/с:", company.get("supplier_account", "")),
        ("МФО:", company.get("supplier_mfo", "")),
        ("Регис. код плател. НДС:", company.get("supplier_vat_code", "")),
    ]
    # Mijoz ma'lumoti: legal_person$export (buyer) dan, bo'lmasa order'dan.
    b = buyer or {}
    buyer_rows = [
        ("Покупатель:", b.get("name") or _g(deal, "person_name")),
        ("Адрес:", b.get("address")
            or _g(deal, "delivery_address_full", "delivery_address_short")),
        ("Тел:", b.get("main_phone") or ""),
        ("ИНН:", b.get("tin") or _g(deal, "person_tin")),
        ("Р/с:", _main_bank(b, "bank_account_code")),
        ("МФО:", _main_bank(b, "mfo")),
        ("Регис. код плател. НДС:", b.get("vat_code") or ""),
    ]
    info_start = 4
    label_align = Alignment(horizontal="left", vertical="center", wrap_text=False)
    for i in range(len(supplier_rows)):
        r = info_start + i
        slabel, sval = supplier_rows[i]
        blabel, bval = buyer_rows[i]
        # Поставщик (chap): yorliq A:B, qiymat C:E
        ws.merge_cells(f"A{r}:B{r}")
        ws[f"A{r}"] = slabel
        ws[f"A{r}"].font = small
        ws[f"A{r}"].alignment = label_align
        ws.merge_cells(f"C{r}:E{r}")
        ws[f"C{r}"] = sval
        ws[f"C{r}"].font = small
        ws[f"C{r}"].alignment = left
        # Покупатель (o'ng): yorliq F:G, qiymat H:J
        ws.merge_cells(f"F{r}:G{r}")
        ws[f"F{r}"] = blabel
        ws[f"F{r}"].font = small
        ws[f"F{r}"].alignment = label_align
        ws.merge_cells(f"H{r}:J{r}")
        ws[f"H{r}"] = bval
        ws[f"H{r}"].font = small
        ws[f"H{r}"].alignment = left
        # qator balandligi — eng uzun qiymatga qarab (yopishib qolmasligi uchun)
        lines = max(_est_lines(sval, 40), _est_lines(bval, 33))
        ws.row_dimensions[r].height = 12 * lines + 3

    # ── Jadval sarlavhasi ──
    hdr = info_start + len(supplier_rows) + 1   # bitta bo'sh qator tashlab
    headers = {
        "A": "№", "D": "ИКПУ", "E": "кол-во", "F": "Цена с НДС",
        "G": "Стоимость\nпоставки", "H": "ставка", "I": "НДС",
        "J": "Стоим. поставки с\nучетом НДС",
    }
    ws.merge_cells(f"B{hdr}:C{hdr}")
    ws[f"B{hdr}"] = "Номенклатура"
    for col, text in headers.items():
        ws[f"{col}{hdr}"] = text
    for col in "ABCDEFGHIJ":
        cell = ws[f"{col}{hdr}"]
        cell.font = bold
        cell.alignment = center
        cell.border = box
    ws.row_dimensions[hdr].height = 28

    # ── Tovarlar (order_products) ──
    items = _g(deal, "order_products", default=[])
    money = "#,##0.00"
    start = hdr + 1
    r = start
    sum_qty = sum_net = sum_vat = sum_total = 0.0
    for idx, it in enumerate(items, start=1):
        name = _clean_product_name(_g(it, "product_name"))
        ikpu = _ikpu_for(name)
        qty = _num(_g(it, "sold_quant", "order_quant", default=0))
        price = _num(_g(it, "product_price", default=0))         # Цена с НДС
        rate = _num(_g(it, "vat_percent", default=12)) / 100.0   # 0.12
        total = _num(_g(it, "sold_amount", default=0)) or (qty * price)  # J
        vat = _num(_g(it, "vat_amount", default=0))
        if not vat and rate:
            vat = total - total / (1 + rate)
        net = total - vat                                        # G

        ws[f"A{r}"] = idx
        ws.merge_cells(f"B{r}:C{r}")
        ws[f"B{r}"] = name
        ws[f"D{r}"] = ikpu
        ws[f"E{r}"] = qty
        ws[f"F{r}"] = price
        ws[f"G{r}"] = round(net, 2)
        ws[f"H{r}"] = rate
        ws[f"I{r}"] = round(vat, 2)
        ws[f"J{r}"] = round(total, 2)

        for col in "ABCDEFGHIJ":
            cell = ws[f"{col}{r}"]
            cell.border = box
            cell.font = reg
            if col in "FGIJ":
                cell.number_format = money
                cell.alignment = right
            elif col == "B":
                cell.alignment = left
            elif col == "H":
                cell.number_format = "0.00"
                cell.alignment = center
            elif col == "E":
                cell.alignment = center
            else:
                cell.alignment = center

        # qator balandligi — uzun mahsulot nomi yopishib qolmasligi uchun
        ws.row_dimensions[r].height = 12 * _est_lines(name, 35) + 3

        sum_qty += qty
        sum_net += net
        sum_vat += vat
        sum_total += total
        r += 1

    # ── Итого ──
    tr = r
    ws.merge_cells(f"B{tr}:C{tr}")
    ws[f"B{tr}"] = "Итого:"
    ws[f"E{tr}"] = sum_qty
    ws[f"G{tr}"] = round(sum_net, 2)
    ws[f"I{tr}"] = round(sum_vat, 2)
    ws[f"J{tr}"] = round(sum_total, 2)
    for col in "ABCDEFGHIJ":
        cell = ws[f"{col}{tr}"]
        cell.border = box
        cell.font = bold
        if col in "GIJ":
            cell.number_format = money
            cell.alignment = right
        elif col == "E":
            cell.alignment = center
        elif col == "B":
            cell.alignment = right

    # ── Summa so'z bilan ──
    wr = tr + 1
    words = _rus_amount_words(sum_total).capitalize()
    ws.merge_cells(f"B{wr}:J{wr}")
    ws[f"B{wr}"] = f"Всего отпущено на сумму: {words} сум 00 тийин"
    ws[f"B{wr}"].font = bold
    ws[f"B{wr}"].alignment = left

    # ── Imzolar ──
    sr = wr + 2
    ws[f"B{sr}"] = "ПОСТАВЩИК"
    ws[f"G{sr}"] = "ПОКУПАТЕЛЬ"
    ws[f"B{sr}"].font = bold
    ws[f"G{sr}"].font = bold
    dr = sr + 1
    ws.merge_cells(f"B{dr}:E{dr}")
    ws[f"B{dr}"] = (f'Директор: {company.get("supplier_director", "")} '
                    f'_____________________\nМ.П')
    ws[f"B{dr}"].font = reg
    ws[f"B{dr}"].alignment = left
    ws.merge_cells(f"G{dr}:J{dr}")
    ws[f"G{dr}"] = "Директор: __________________\nМ.П"
    ws[f"G{dr}"].font = reg
    ws[f"G{dr}"].alignment = left
    ws.row_dimensions[dr].height = 28

    ws.page_setup.orientation = "landscape"
    ws.print_options.horizontalCentered = True

    deal_id = _g(deal, "deal_id", "id", default="deal")
    path = os.path.join(tempfile.gettempdir(),
                        f"Spetsifikatsiya_{company['name']}_{deal_id}.xlsx")
    wb.save(path)
    return path
